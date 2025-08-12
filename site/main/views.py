from django.shortcuts import render, redirect

from django.core.exceptions import ValidationError
from .forms import RegistrationForm, LoginForm
from .forms import CriterionForm
from django.conf import settings
from django.db import connection
from django.core.cache import cache
from django.utils import timezone
from .models import GeneratedDocument
from .models import UserAvatar
import os, bcrypt,re
from django_ratelimit.decorators import ratelimit
from django.core.validators import RegexValidator
from django.contrib.auth.password_validation import validate_password
from django.contrib import messages
from django.http import JsonResponse
from .models import City

from datetime import date
from docxtpl import DocxTemplate
from urllib.parse import quote


from django.http import HttpResponse
from docxtpl import DocxTemplate
from io import BytesIO
from django.db import connection

from docx import Document


def index(request):
    return render(request, 'main/index.html')


def password_check(request,password,loggin,spesial_password,id,is_admin):

        lock_key = f'login_lock_{loggin}'
        lock_time = cache.get(lock_key)
        print(f"Lock time from cache: {lock_time}") 

        if lock_time:
            remaining = (lock_time - timezone.now()).seconds
            return render(request, 'main/about.html', {
                'form': RegistrationForm(),
                'error_message': f"Потворите попыптку через {remaining} секунд"
            })

        fail_count = cache.get(f'login_fails_{loggin}', 0)

        try:
            hash_bytes = spesial_password.encode('utf-8')
                
            if bcrypt.checkpw(password.encode('utf-8'), hash_bytes):
                    success_message = f"Авторизация для {loggin} выполнена успешно"

                    view_name = parsing(loggin)

                    if is_admin:
                        return redirect('success1', username = view_name,user_id = id)
                    
                    else:
                        return redirect('success', username = view_name,user_id = id)


            else:
                    
                    fail_count += 1
                    cache.set(f'login_fails_{loggin}', fail_count, timeout=300)

                    if fail_count > 4:
                        cache.set(lock_key, timezone.now() + timezone.timedelta(minutes=2), timeout=120)

                    error_message = 'Неверный пароль'
                    return render(request, 'main/about.html',
                            {'form': RegistrationForm(), 'error_message': error_message})
                
        except Exception as e:
                error_message = 'Ошибка аутентификации'
                return render(request, 'main/about.html',
                           {'form': RegistrationForm(), 'error_message': error_message})


def validation(username):

    validator = RegexValidator(
    r'^[a-zA-Zа-яА-ЯёЁ0-9_@.+\-]+$',
    'Invalid username!'
)
    try:
        validator(username)
        return None
    except ValidationError:
        return 'Username contains forbidden characters!'
    
    
# @ratelimit(key='ip', rate='8/m',block=True)
def about(request):
    """Авторизация"""
    if request.method == 'POST':
       
        username = request.POST.get('username')
        password = request.POST.get('password')
        ip = request.META.get('REMOTE_ADDR')    


        ip_fail_key = f'login_fails_ip_{ip}'
        ip_lock_key = f'login_lock_ip_{ip}'

        ip_lock = cache.get(ip_lock_key)

        if ip_lock:
            remaining = (ip_lock - timezone.now()).seconds
            return render(request, 'main/about.html', {
                'form': RegistrationForm(),
                'error_message': f"Потворите попыптку через {remaining} секунд"
            })

        error_message = validation(username)
        if error_message:
            return render(request, 'main/about.html', {
                'form': RegistrationForm(),
                'error_message': error_message
            })

        with connection.cursor() as cursor:
            # сначала ищем спец пользователей
            cursor.execute(
                "SELECT loggin,password,id FROM loging_password WHERE loggin = %s",
                [username]
            )
            special = cursor.fetchone()

            if special:
                return password_check(request, password,special[0],special[1],special[2],True)

            # тут обычных
            cursor.execute(
                "SELECT user_password,name,id FROM users WHERE name = %s",
                [username]
            )
            user_data = cursor.fetchone()

            if user_data:
                return password_check(request, password,user_data[1],user_data[0],user_data[2],False)
            
            else:
                ip_fails = cache.get(ip_fail_key, 0) + 1
                cache.set(ip_fail_key, ip_fails, timeout=300)

                if ip_fails > 5:
                    cache.set(ip_lock_key, 
                            timezone.now() + timezone.timedelta(minutes=2), 
                            timeout=120)
                    
                error_message = 'Пользователь не найден'
                return render(request, 'main/about.html', 
                           {'form': RegistrationForm(), 'error_message': error_message})

    else:
        return render(request, 'main/about.html', {'form': RegistrationForm()})

def registr(request):
    error_message = None
    
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        confirm_password = request.POST.get('confirm_password')

        with connection.cursor() as cursor:
                cursor.execute(
                    "SELECT name FROM users WHERE name = %s",
                    [username]
                )
                not_uniq = cursor.fetchone()

                print(not_uniq)

                if not_uniq:
                    error_message = 'Логин занят'
                    return render(request, 'main/registr.html', 
                            {'error_message': error_message})

        if password != confirm_password:
            error_message = 'Пароли не совпадают'
            return render(request, 'main/registr.html', 
                        {'error_message': error_message})
        
        error_message = validation(username)
        if error_message:
            return render(request, 'main/registr.html', {
                'form': RegistrationForm(),
                'error_message': error_message
            })

        try:
            validate_password(password,username)
        except ValidationError as ve:
            error_message = 'Ошибка в пароле: ' + ', '.join(ve.messages)
            return render(request,'main/registr.html',
                        {'error_message':error_message})

        try:
            salt = bcrypt.gensalt()
            hashed_password = bcrypt.hashpw(password.encode('utf-8'), salt)
            hashed_password_str = hashed_password.decode('utf-8')


            with connection.cursor() as cursor:
                cursor.execute(
                    "INSERT INTO users (name, user_password) VALUES (%s, %s)",
                    [username, hashed_password_str]
                )
                

            with connection.cursor() as cursor:
                cursor.execute(
                    "SELECT id FROM users WHERE name = %s",
                    [username]
                )
                user_id = cursor.fetchone()

            view_name = parsing(username)
            print(parsing(username))
            return redirect('success', username=view_name,user_id = user_id[0])
        
        except Exception as e:
            error_message = f"Ошибка при регистрации: {str(e)}"
    
    return render(request, 'main/registr.html', {'error_message': error_message})


def parsing(username):

    if "@" in username:
        dog_sep = username.split("@")[0]
        dot_sep = username.split(".")[0]

        return dog_sep if len(dog_sep) < len(dot_sep) else dot_sep
    
    else:
         return username
 
def city_autocomplete(request):

    q = request.GET.get('q','')
    print("dsdsadsadas")
    if q:
        cities = City.objects.filter(name__icontains=q).order_by('name')[:10]
        results = list(cities.values_list('name', flat=True))
    else:
        results = []
    return JsonResponse(results, safe=False)


def check_personal_data(user_id,request):


    with connection.cursor() as cursor:
        cursor.execute("SELECT * FROM personal_data WHERE user_id = %s", [user_id])
        is_data = cursor.fetchone()

        if not is_data:
            messages.error(request, "Нет данных для генерации документа.")
            return redirect('success', username=request.user.username, user_id=user_id)
        
    name = is_data[1]
    surname = is_data[2]
    last_name = is_data[3]
    dob = is_data[4]
    city = is_data[5]
    current_time = timezone.now()
    current_date = current_time.date().strftime('%d.%m.%Y')

    today = timezone.now().date()

    if dob:
        dob_str = dob.strftime('%d.%m.%Y')
    else:
        dob_str = ''

    age = today.year - dob.year - ((today.month, today.day) < (dob.month, dob.day))

    of_legal_age = "Да"

    if age < 18: 
        of_legal_age = "Нет"


    context = {
        'ФИО': f'{name} {surname} {last_name}',
        'Город_проживания': city,
        'Дата_рождения': dob_str,
        'дата_выдачи': current_date,
        'Серия':'Не указанно',
        'Номер':'Не указанно',
        'Кем':'Не указанно',
        'Когда':'Не указанно',
        'Код':'Не указанно',
        'Возраст':age,
        'Совершеннолетний':of_legal_age
        }

    return context

def lack_data(user_id,request):
    
    context = check_personal_data(user_id,request)

    template_file = request.POST.get('template')
    path = f"main/templates/documents/{template_file}"
        
    doc = Document(path)
    text = "\n".join([p.text for p in doc.paragraphs])
    fields = re.findall(r'\{\{\s*([^}]+)\s*\}\}', text)

    not_match = []
    keys = [i for i in context.keys()]
    values = [i for i in context.values()]

    for field in fields:
        if field in context:
            if context[field] == "Не указанно" or context[field] == '' or context[field] is None:
                not_match.append(field)
        else:
            not_match.append(field)

    return not_match


def generate_doc(user_id, request,surname):

    template_file = request.POST.get('template')
    path = f"main/templates/documents/{template_file}"

    context = check_personal_data(user_id,request)

    doc = Document(path)
    text = "\n".join([p.text for p in doc.paragraphs])
    fields = re.findall(r'\{\{\s*([^}]+)\s*\}\}', text)
    not_match = []

    keys = [i for i in context.keys()]
    for i in fields:
        if i not in keys:
            not_match.append(i)

    doc = DocxTemplate(path)
    file_stream = BytesIO()
    doc.render(context)
    doc.save(file_stream)
    file_stream.seek(0)

    filename = template_file
    filename = filename.split(".docx")[0]
    print(filename)
    quoted_filename = quote(filename)
    quoted_surname = quote(surname)
    
    response = HttpResponse(
            file_stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{quoted_filename}_{quoted_surname}.docx"

    GeneratedDocument.objects.create(
        user_id = user_id,
        template_name = template_file,
        file = template_file
    )

    return response
    


def get_client_ip(request):
   
    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded_for:
        ip = x_forwarded_for.split(',')[0].strip()
    else:
        ip = request.META.get('REMOTE_ADDR')
    return ip

def write_accept(user_id,request):


    ip_address = get_client_ip(request)
    user_agent = request.META.get('HTTP_USER_AGENT', '')

    with connection.cursor() as cursor:
        cursor.execute(
        "SELECT * FROM personal_data_agreement WHERE user_id = %s",
        [user_id]
        )

        is_data = cursor.fetchone()

    if not is_data:
        with connection.cursor() as cursor:
            cursor.execute(
            "INSERT INTO personal_data_agreement (user_id,agreed_at,ip_address, user_agent) VALUES (%s,NOW(),%s,%s) ",
            [user_id,ip_address,user_agent]
            )

def check_accept(user_id):

    with connection.cursor() as cursor:
        cursor.execute(
        "SELECT * FROM personal_data_agreement WHERE user_id = %s",
        [user_id]
        )
        is_data = cursor.fetchone()

    if is_data:
        return True
    return False
    


def success(request, username,user_id):
    
    form_data = {
            'first_name': '',
            'last_name': '',
            'middle_name': '',
            'birth_date': '',
            'city': ''
        }
    
    add = []

    history = GeneratedDocument.objects.filter(user_id=user_id).order_by('-created_at')
    
    folder_path = os.path.join('main', 'templates', 'documents')
    files = os.listdir(folder_path)
    templates = [f for f in files if f.endswith('.docx')]
    
    accept_given = check_accept(user_id)
    print(accept_given,"Accept")

    avatar_file = request.FILES.get('avatar')

    info_personal = get_data(user_id)
    surname = info_personal['last_name']

    if request.method == 'POST':
        print(request.POST)

        if avatar_file:
            avatar_obj, created = UserAvatar.objects.get_or_create(user_id=user_id)
            avatar_obj.avatar = avatar_file
            avatar_obj.save()   

        avatar_url = ''
        avatar = UserAvatar.objects.filter(user_id=user_id).first()
        if avatar and avatar.avatar:
            avatar_url = avatar.avatar.url

        if 'clear_history' in request.POST:
            return delete_history(request,user_id,username)

        if 'generate_doc_with_missing' in request.POST:
            template_file = request.POST.get('template')
            print(template_file,"ФАЙЛ")
            if not template_file:
                messages.error(request, "Не выбран шаблон документа.")
                return redirect('success', username=username, user_id=user_id)
            
            context = check_personal_data(user_id, request)
            for key in request.POST:
                if key.startswith('missing_'):
                    field_name = key[len('missing_'):]
                    context[field_name] = request.POST[key].strip()

            return generate_doc_with_context(template_file, context,user_id,surname)

        # Запись согласия
        if 'accept-consent' in request.POST:

            write_accept(user_id, request)
            messages.success(request, 'Согласие сохранено')
            return redirect('success', username=username, user_id=user_id) 

        # создание доков
        if 'generate_doc_btn' in request.POST:
            
            # проверка на заполненность доков
            if 'download_doc' in request.POST:
                print("ЕЕЕБОЙ")
                add = lack_data(user_id,request)
                print(add,"Не хвататет")
                if add:
                    
                    template_file = request.POST.get('template')
                    request.session['template_file'] = template_file
                    request.session['lack_data'] = add
                    request.session['form_data'] = info_personal
                    return redirect('success', username=username, user_id=user_id)
                
                
            
            return generate_doc(user_id, request,surname)


        
    
        last_name = request.POST.get('last_name', '').strip()
        first_name = request.POST.get('first_name', '').strip()
        middle_name = request.POST.get('middle_name', '').strip()
        birth_date = request.POST.get('birth_date', '').strip()
        city = request.POST.get('city', '').strip()

                

        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT * FROM personal_data WHERE user_id = %s",
                [user_id]
                )
            is_data = cursor.fetchone()

        if is_data:
                with connection.cursor() as cursor:
                    cursor.execute("""
                        UPDATE personal_data SET name = %s, surname = %s, last_name = %s, dob = %s, sity = %s 
                        WHERE user_id = %s""", [first_name, last_name, middle_name, birth_date, city, user_id])

        # сохранение данных      
        else:
            with connection.cursor() as cursor:
                cursor.execute(
                    "INSERT INTO personal_data (user_id, name,surname,last_name,dob,sity) VALUES (%s, %s,%s,%s,%s,%s)",
                    [user_id,last_name,first_name,middle_name,birth_date,city]
                    )
                
        form_data = {
            'first_name': first_name,
            'last_name': last_name,
            'middle_name': middle_name,
            'birth_date': birth_date,
            'city': city
        }
        
        # lack_data1 = []
        messages.success(request, 'Профиль успешно обновлён!')

        
    # GET запрос
    else:
        
        form_data = get_data(user_id)

        avatar_url = ''
        avatar = UserAvatar.objects.filter(user_id=user_id).first()
        if avatar and avatar.avatar:
            avatar_url = avatar.avatar.url

    lack_data1 = request.session.pop('lack_data', [])
    if not lack_data1:
        lack_data1 = []
    form_data_from_session = request.session.pop('form_data', None)
    if form_data_from_session:
        form_data = form_data_from_session

    template_file = request.session.pop('template_file', '')

    return render(request, 'main/success.html', {
        'username': username,
        'user_id': user_id,
        'form_data': form_data,
        'full_name': f"{form_data['last_name']} {form_data['first_name']} {form_data['middle_name']}",
        'birth_date': form_data['birth_date'],
        'city': form_data['city'],
        'avatar_url': avatar_url,
        'show_consent_modal': not accept_given,
        'templates': templates,
        'lack_data':lack_data1,
        'template_file': template_file,
        'history': history
    })


def get_data(user_id):

    form_data = {
        'first_name': '',
        'last_name': '',
        'middle_name': '',
        'birth_date': '',
        'city': ''
    }

    with connection.cursor() as cursor:
        cursor.execute("SELECT * FROM personal_data WHERE user_id = %s", [user_id])
        row = cursor.fetchone()
        if row:
            birth_date = row[4].strftime('%Y-%m-%d') if row[4] else ''

            form_data = {
                'first_name': row[1],
                'last_name': row[2],
                'middle_name': row[3],
                'birth_date': birth_date,
                'city': row[5]
                }
            
    return form_data

def generate_doc_with_context(template_file, context,user_id,surname):
    path = f"main/templates/documents/{template_file}"
    doc = DocxTemplate(path)
    file_stream = BytesIO()
    doc.render(context)
    doc.save(file_stream)
    file_stream.seek(0)

    filename = template_file
    filename = filename.split(".docx")[0]
    print(filename)
    quoted_filename = quote(filename)
    quoted_surname = quote(surname)

    response = HttpResponse(
        file_stream.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{quoted_filename}_{quoted_surname}.docx"


    GeneratedDocument.objects.create(
        user_id = user_id,
        template_name = template_file,
        file = template_file
    )

    return response


def delete_history(request,user_id,username):
    GeneratedDocument.objects.filter(user_id=user_id).delete()
    messages.success(request, 'История документов успешно очищена.')
    return redirect('success', username=username, user_id=user_id)
    

def success1(request,username,user_id):
    if request.method == 'POST':
        last_name = request.POST.get('last_name', '').strip()
        first_name = request.POST.get('first_name', '').strip()
        middle_name = request.POST.get('middle_name', '').strip()
        birth_date = request.POST.get('birth_date', '').strip()
        city = request.POST.get('city', '').strip()

        

        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO personal_data (user_id, name,surname,last_name,dob,sity) VALUES (%s, %s,%s,%s,%s,%s)",
                [user_id,last_name,first_name,middle_name,birth_date,city]
            )

        
    return render(request, 'main/success.html',{'username':username},{'user_id':user_id})

        

def criterion_view(request):
    if request.method == 'POST':
        form = CriterionForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            return redirect('success')
    else:
        form = CriterionForm()
    return render(request, 'main/criterion.html', {'form': form})

def media(request):
    media_root = settings.MEDIA_ROOT
    files = os.listdir(media_root)
    return render(request, 'main/media.html', {'files': files})